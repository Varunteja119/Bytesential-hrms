from app.database.models.employee import EmploymentStatus
from app.database.models.rbac import Permission, Role
from app.database.models.recruitment import Candidate, CandidateStatus, Job, JobStatus
from app.database.models.user import User
from app.security.password import hash_password


def _make_hr_user(db_session) -> None:
    codes = ["employee:read", "employee:write"]
    perms = [Permission(code=c) for c in codes]
    db_session.add_all(perms)
    role = Role(name="hr", permissions=perms)
    db_session.add(role)
    user = User(email="hr@example.com", hashed_password=hash_password("supersecret1"), full_name="HR Person", roles=[role])
    db_session.add(user)
    employee_role = Role(name="employee", permissions=[])
    db_session.add(employee_role)
    db_session.commit()


def _hr_headers(client) -> dict:
    resp = client.post("/api/v1/auth/login", json={"email": "hr@example.com", "password": "supersecret1"})
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def _make_accepted_candidate(db_session, email="newhire@example.com") -> str:
    hr = db_session.query(User).filter(User.email == "hr@example.com").first()
    job = Job(title="Engineer", department="Eng", description="desc", status=JobStatus.OPEN, created_by_id=hr.id)
    db_session.add(job)
    db_session.commit()
    candidate = Candidate(full_name="New Hire", email=email, job_id=job.id, status=CandidateStatus.ACCEPTED)
    db_session.add(candidate)
    db_session.commit()
    return str(candidate.id)


def _provision(client, headers, candidate_id):
    return client.post(
        "/api/v1/employees/provision",
        json={"candidate_id": candidate_id, "department": "Engineering", "designation": "Software Engineer", "date_of_joining": "2026-08-01"},
        headers=headers,
    )


def _complete_profile(client, employee_headers):
    return client.patch(
        "/api/v1/employees/me",
        json={"phone": "9999999999", "address": "123 Main St", "date_of_birth": "1995-01-01", "bank_account_number": "123456"},
        headers=employee_headers,
    )


def test_new_hire_must_change_password_flag_set(client, db_session):
    _make_hr_user(db_session)
    headers = _hr_headers(client)
    candidate_id = _make_accepted_candidate(db_session)
    temp_password = _provision(client, headers, candidate_id).json()["temp_password"]

    login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": temp_password})
    employee_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    me = client.get("/api/v1/auth/me", headers=employee_headers)
    assert me.json()["must_change_password"] is True


def test_change_password_wrong_current_password_rejected(client, db_session):
    _make_hr_user(db_session)
    headers = _hr_headers(client)
    candidate_id = _make_accepted_candidate(db_session)
    temp_password = _provision(client, headers, candidate_id).json()["temp_password"]

    login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": temp_password})
    employee_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    resp = client.post(
        "/api/v1/auth/change-password",
        json={"current_password": "wrongpassword", "new_password": "brandnewpass123"},
        headers=employee_headers,
    )
    assert resp.status_code == 401


def test_change_password_clears_flag_and_new_password_works(client, db_session):
    _make_hr_user(db_session)
    headers = _hr_headers(client)
    candidate_id = _make_accepted_candidate(db_session)
    temp_password = _provision(client, headers, candidate_id).json()["temp_password"]

    login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": temp_password})
    employee_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    resp = client.post(
        "/api/v1/auth/change-password",
        json={"current_password": temp_password, "new_password": "brandnewpass123"},
        headers=employee_headers,
    )
    assert resp.status_code == 200
    assert resp.json()["must_change_password"] is False

    # old temp password no longer works
    old_login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": temp_password})
    assert old_login.status_code == 401

    # new password works
    new_login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": "brandnewpass123"})
    assert new_login.status_code == 200


def test_activation_blocked_if_password_not_changed_even_with_docs_and_profile(client, db_session):
    import io
    from docx import Document
    from app.main import app
    from app.services.storage import get_storage_client

    class FakeStorage:
        def __init__(self):
            self.files = {}

        def upload(self, key, data, content_type):
            self.files[key] = data

        def download(self, key):
            return self.files[key]

    app.dependency_overrides[get_storage_client] = lambda: FakeStorage()
    try:
        _make_hr_user(db_session)
        headers = _hr_headers(client)
        candidate_id = _make_accepted_candidate(db_session)
        provision = _provision(client, headers, candidate_id).json()
        employee_id = provision["employee"]["id"]
        temp_password = provision["temp_password"]

        login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": temp_password})
        employee_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

        # Deliberately SKIP change-password. Complete everything else.
        _complete_profile(client, employee_headers)
        for doc_type in ["aadhaar", "pan", "bank_proof"]:
            doc = Document()
            doc.add_paragraph(f"Fake {doc_type}")
            buf = io.BytesIO()
            doc.save(buf)
            resp = client.post(
                f"/api/v1/employees/{employee_id}/documents?document_type={doc_type}",
                files={"file": (f"{doc_type}.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                headers=employee_headers,
            )
            client.patch(f"/api/v1/employees/{employee_id}/documents/{resp.json()['id']}/verify", headers=headers)

        # profile complete + all required docs verified, but password never changed -> still blocked
        resp = client.post(f"/api/v1/employees/{employee_id}/activate", headers=headers)
        assert resp.status_code == 400
        assert "password" in resp.json()["error"]["message"].lower()
    finally:
        del app.dependency_overrides[get_storage_client]


def test_activation_requires_specific_required_document_types(client, db_session):
    import io
    from docx import Document
    from app.main import app
    from app.services.storage import get_storage_client

    class FakeStorage:
        def __init__(self):
            self.files = {}

        def upload(self, key, data, content_type):
            self.files[key] = data

        def download(self, key):
            return self.files[key]

    app.dependency_overrides[get_storage_client] = lambda: FakeStorage()
    try:
        _make_hr_user(db_session)
        headers = _hr_headers(client)
        candidate_id = _make_accepted_candidate(db_session)
        provision = _provision(client, headers, candidate_id).json()
        employee_id = provision["employee"]["id"]
        temp_password = provision["temp_password"]

        login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": temp_password})
        employee_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

        # change password
        client.post(
            "/api/v1/auth/change-password",
            json={"current_password": temp_password, "new_password": "brandnewpass123"},
            headers=employee_headers,
        )
        # complete profile
        _complete_profile(client, employee_headers)

        # activation still blocked -- ZERO documents uploaded (this is exactly the gap that used to slip through)
        blocked = client.post(f"/api/v1/employees/{employee_id}/activate", headers=headers)
        assert blocked.status_code == 400
        assert "document" in blocked.json()["error"]["message"].lower()

        def upload_doc(doc_type: str):
            doc = Document()
            doc.add_paragraph(f"Fake {doc_type} content")
            buf = io.BytesIO()
            doc.save(buf)
            resp = client.post(
                f"/api/v1/employees/{employee_id}/documents?document_type={doc_type}",
                files={"file": (f"{doc_type}.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                headers=employee_headers,
            )
            assert resp.status_code == 201
            return resp.json()["id"]

        # upload only ONE of the three required types, verify it -- still blocked (2 more required)
        aadhaar_id = upload_doc("aadhaar")
        client.patch(f"/api/v1/employees/{employee_id}/documents/{aadhaar_id}/verify", headers=headers)
        still_blocked = client.post(f"/api/v1/employees/{employee_id}/activate", headers=headers)
        assert still_blocked.status_code == 400

        # upload and verify the remaining required types
        pan_id = upload_doc("pan")
        client.patch(f"/api/v1/employees/{employee_id}/documents/{pan_id}/verify", headers=headers)
        bank_id = upload_doc("bank_proof")
        client.patch(f"/api/v1/employees/{employee_id}/documents/{bank_id}/verify", headers=headers)

        # a non-required doc type also present, doesn't matter either way
        upload_doc("other")

        activate = client.post(f"/api/v1/employees/{employee_id}/activate", headers=headers)
        assert activate.status_code == 200
        assert activate.json()["employment_status"] == "active"
    finally:
        del app.dependency_overrides[get_storage_client]


def test_onboarding_status_checklist_reflects_progress(client, db_session):
    _make_hr_user(db_session)
    headers = _hr_headers(client)
    candidate_id = _make_accepted_candidate(db_session)
    provision = _provision(client, headers, candidate_id).json()
    employee_id = provision["employee"]["id"]
    temp_password = provision["temp_password"]

    # HR view before anything happens
    status_resp = client.get(f"/api/v1/employees/{employee_id}/onboarding-status", headers=headers)
    assert status_resp.status_code == 200
    status = status_resp.json()
    assert status["credentials_generated"] is True
    assert status["password_changed"] is False
    assert status["profile_completed"] is False
    assert status["ready_for_activation"] is False
    assert len(status["documents"]) == 3  # aadhaar, pan, bank_proof
    assert all(not d["uploaded"] for d in status["documents"])

    login = client.post("/api/v1/auth/login", json={"email": "newhire@example.com", "password": temp_password})
    employee_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    client.post(
        "/api/v1/auth/change-password",
        json={"current_password": temp_password, "new_password": "brandnewpass123"},
        headers=employee_headers,
    )
    _complete_profile(client, employee_headers)

    self_status = client.get("/api/v1/employees/me/onboarding-status", headers=employee_headers)
    assert self_status.status_code == 200
    body = self_status.json()
    assert body["password_changed"] is True
    assert body["profile_completed"] is True
    assert body["ready_for_activation"] is False  # still missing documents


def test_employee_cannot_view_another_employees_onboarding_status(client, db_session):
    _make_hr_user(db_session)
    headers = _hr_headers(client)

    c1 = _make_accepted_candidate(db_session, email="hire1@example.com")
    e1 = _provision(client, headers, c1).json()
    c2 = _make_accepted_candidate(db_session, email="hire2@example.com")
    e2 = _provision(client, headers, c2).json()

    login1 = client.post("/api/v1/auth/login", json={"email": "hire1@example.com", "password": e1["temp_password"]})
    headers1 = {"Authorization": f"Bearer {login1.json()['access_token']}"}

    resp = client.get(f"/api/v1/employees/{e2['employee']['id']}/onboarding-status", headers=headers1)
    assert resp.status_code == 403
